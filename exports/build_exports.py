#!/usr/bin/env python3
"""Convert the Markdown documents to DOCX and the slide outlines to PPTX, using the Vytalix brand.
Run: python3 exports/build_exports.py  (outputs into exports/docx and exports/pptx)"""
import os,re,glob
from docx import Document
from docx.shared import Pt,RGBColor,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PIn,Pt as PPt
from pptx.dml.color import RGBColor as PRGB
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_D=os.path.join(ROOT,'exports','docx'); OUT_P=os.path.join(ROOT,'exports','pptx')
os.makedirs(OUT_D,exist_ok=True); os.makedirs(OUT_P,exist_ok=True)
BLUE=RGBColor(0x15,0x38,0x92); INK=RGBColor(0x10,0x19,0x2B); SLATE=RGBColor(0x4B,0x55,0x68); ROSE=RGBColor(0xC9,0x3D,0x63)
LOGO=os.path.join(ROOT,'website','assets','img','vytalix-logo.png')

def inline(par,text):
    """Bold/italic/code inline markdown -> runs."""
    text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',text)
    for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)',text):
        if not tok: continue
        if tok.startswith('**'): r=par.add_run(tok[2:-2]); r.bold=True
        elif tok.startswith('`'): r=par.add_run(tok[1:-1]); r.font.name='Consolas'; r.font.size=Pt(9)
        elif tok.startswith('*') and len(tok)>2: r=par.add_run(tok[1:-1]); r.italic=True
        else: par.add_run(tok)

def md_to_docx(src,dst):
    doc=Document()
    st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5)
    for s in ('Heading 1','Heading 2','Heading 3'):
        doc.styles[s].font.color.rgb=BLUE; doc.styles[s].font.name='Cambria'
    sec=doc.sections[0]; sec.left_margin=sec.right_margin=Inches(0.9); sec.top_margin=sec.bottom_margin=Inches(0.8)
    if os.path.exists(LOGO): doc.add_picture(LOGO,width=Inches(1.6))
    lines=open(src,encoding='utf-8').read().split('\n'); i=0; table=[]
    def flush_table():
        nonlocal table
        rows=[r for r in table if not re.match(r'^\|?\s*:?-{2,}',r)]
        if not rows: table=[]; return
        cells=[[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
        ncol=max(len(c) for c in cells)
        t=doc.add_table(rows=len(cells),cols=ncol); t.style='Light Grid Accent 1'
        for ri,row in enumerate(cells):
            for ci in range(ncol):
                cell=t.cell(ri,ci); cell.text=''; p=cell.paragraphs[0]; inline(p,row[ci] if ci<len(row) else '')
                for r in p.runs: r.font.size=Pt(8.5); r.bold=r.bold or ri==0
        doc.add_paragraph(); table=[]
    while i<len(lines):
        l=lines[i].rstrip()
        if l.startswith('|'): table.append(l); i+=1; continue
        if table: flush_table()
        if l.startswith('```'):
            i+=1; buf=[]
            while i<len(lines) and not lines[i].startswith('```'): buf.append(lines[i]); i+=1
            p=doc.add_paragraph(); r=p.add_run('\n'.join(buf)); r.font.name='Consolas'; r.font.size=Pt(8.5); i+=1; continue
        m=re.match(r'^(#{1,6})\s+(.*)',l)
        if m:
            lvl=min(len(m.group(1)),3); h=doc.add_heading(re.sub(r'[*`]','',m.group(2)),level=lvl if lvl>1 else 0 if i<3 else 1)
            i+=1; continue
        if re.match(r'^\s*[-*]\s+',l):
            p=doc.add_paragraph(style='List Bullet'); inline(p,re.sub(r'^\s*[-*]\s+','',l)); i+=1; continue
        if re.match(r'^\s*\d+\.\s+',l):
            p=doc.add_paragraph(style='List Number'); inline(p,re.sub(r'^\s*\d+\.\s+','',l)); i+=1; continue
        if l.startswith('>'):
            p=doc.add_paragraph(); inline(p,l.lstrip('> ')); p.paragraph_format.left_indent=Inches(0.3)
            for r in p.runs: r.font.color.rgb=SLATE; r.italic=True
            i+=1; continue
        if l.strip() in ('---','***'): i+=1; continue
        if l.strip(): p=doc.add_paragraph(); inline(p,l)
        i+=1
    if table: flush_table()
    f=doc.sections[0].footer.paragraphs[0]; f.text='Vytalix · Technology for Life · Confidential draft · Every figure is an ESTIMATE unless labelled VERIFIED (see FACTS_BASE.md)'; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in f.runs: r.font.size=Pt(8); r.font.color.rgb=SLATE
    doc.save(dst)

def md_to_pptx(src,dst,title):
    prs=Presentation(); prs.slide_width=PIn(13.333); prs.slide_height=PIn(7.5)
    txt=open(src,encoding='utf-8').read()
    # slides: split on headings starting with '## Slide' or '## ' followed by number
    parts=re.split(r'\n(?=##+\s+(?:Slide\s*)?\d+)',txt)
    slides=[p for p in parts if re.match(r'##+\s+(?:Slide\s*)?\d+',p.strip())]
    def add_slide(head,body_lines,notes):
        s=prs.slides.add_slide(prs.slide_layouts[6])
        bar=s.shapes.add_shape(1,0,0,prs.slide_width,PIn(0.18)); bar.fill.solid(); bar.fill.fore_color.rgb=PRGB(0x15,0x38,0x92); bar.line.fill.background()
        if os.path.exists(LOGO): s.shapes.add_picture(LOGO,PIn(11.4),PIn(0.35),width=PIn(1.6))
        tb=s.shapes.add_textbox(PIn(0.6),PIn(0.5),PIn(10.5),PIn(1.2)); tf=tb.text_frame; tf.word_wrap=True
        p=tf.paragraphs[0]; p.text=head; p.font.size=PPt(30); p.font.bold=True; p.font.color.rgb=PRGB(0x15,0x38,0x92); p.font.name='Cambria'
        bb=s.shapes.add_textbox(PIn(0.6),PIn(1.8),PIn(12.1),PIn(5.2)); bf=bb.text_frame; bf.word_wrap=True
        first=True
        for ln in body_lines:
            q=bf.paragraphs[0] if first else bf.add_paragraph(); first=False
            q.text=ln; q.font.size=PPt(16) if len(ln)<140 else PPt(13); q.font.color.rgb=PRGB(0x10,0x19,0x2B); q.font.name='Calibri'; q.space_after=PPt(6)
        if notes: s.notes_slide.notes_text_frame.text=notes
        ft=s.shapes.add_textbox(PIn(0.6),PIn(7.0),PIn(12),PIn(0.4)); fp=ft.text_frame.paragraphs[0]; fp.text='Vytalix · Technology for Life · Confidential · In development; no claims of approval, validation or traction beyond those stated'; fp.font.size=PPt(9); fp.font.color.rgb=PRGB(0x4B,0x55,0x68)
    # title slide
    s=prs.slides.add_slide(prs.slide_layouts[6]); bg=s.shapes.add_shape(1,0,0,prs.slide_width,prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb=PRGB(0x0F,0x23,0x54); bg.line.fill.background()
    tb=s.shapes.add_textbox(PIn(0.8),PIn(2.6),PIn(11.5),PIn(2.5)); tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=title; p.font.size=PPt(40); p.font.bold=True; p.font.color.rgb=PRGB(255,255,255); p.font.name='Cambria'
    q=tf.add_paragraph(); q.text='Vytalix · Technology for Life · September 2026 · Draft for discussion'; q.font.size=PPt(18); q.font.color.rgb=PRGB(0xC7,0xD0,0xE8)
    for sl in slides:
        lines=[x.rstrip() for x in sl.strip().split('\n')]
        head=re.sub(r'^#+\s+(?:Slide\s*)?\d+\s*[—:\-–.]*\s*','',lines[0]).strip().strip('*')
        body=[]; notes=[]; mode='body'
        for ln in lines[1:]:
            if not ln.strip(): continue
            low=ln.lower()
            if low.startswith(('**notes','**speaker','speaker notes','notes:','**visual','visual:')):
                mode='notes' if 'note' in low else 'visual'; ln=re.sub(r'^\**(speaker notes|notes|visual)\**\s*[:：]?\s*\**','',ln,flags=re.I)
            if mode=='body':
                if low.startswith(('**headline','headline:')): body.insert(0,re.sub(r'^\**headline\**\s*[:：]?\s*','',ln,flags=re.I).strip('*')); continue
                if low.startswith(('**body','body:')): ln=re.sub(r'^\**body\**\s*[:：]?\s*','',ln,flags=re.I)
                for piece in re.split(r'\s+·\s+',ln):
                    piece=re.sub(r'^[-*]\s+','',piece).replace('**','').strip()
                    if piece: body.append(('• '+piece) if not piece.endswith(':') else piece)
            elif mode=='notes': notes.append(ln.replace('**',''))
            else: notes.append('Visual: '+ln.replace('**',''))
        add_slide(head,body[:12],'\n'.join(notes))
    prs.save(dst)

if __name__=='__main__':
    n=0
    for src in sorted(glob.glob(os.path.join(ROOT,'docs','*.md'))+glob.glob(os.path.join(ROOT,'product','*.md'))+glob.glob(os.path.join(ROOT,'investors','*.md'))+[os.path.join(ROOT,'FACTS_BASE.md')]+[p for p in glob.glob(os.path.join(ROOT,'assets','*.md')) if 'DECK' not in p and 'PRESENTATION' not in p]):
        dst=os.path.join(OUT_D,os.path.splitext(os.path.basename(src))[0]+'.docx'); md_to_docx(src,dst); n+=1
    for name,title in [('INVESTOR_PITCH_DECK','Vytalix — Investor pitch'),('CORPORATE_PRESENTATION','Vytalix — Corporate presentation'),('SALES_DECK','Vytalix Advisory & Consulting — Sales deck'),('PARTNERSHIP_DECK','MaternaLink — Partnership deck')]:
        src=os.path.join(ROOT,'assets',name+'.md')
        if os.path.exists(src): md_to_pptx(src,os.path.join(OUT_P,name+'.pptx'),title); n+=1
    print('exports written:',n)
