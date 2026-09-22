#!/usr/bin/env python3
"""Branded Excel workbooks: adds a cover sheet with logo to existing workbooks and builds new working tools.
Run: python3 exports/build_xlsx.py   Outputs: exports/xlsx/*.xlsx"""
import os,csv,glob,shutil
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.drawing.image import Image as XImage
from openpyxl.utils import get_column_letter
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT=os.path.join(ROOT,'exports','xlsx'); os.makedirs(OUT,exist_ok=True)
LOGO=os.path.join(ROOT,'website','assets','img','vytalix-logo.png')
BLUE='153892'; ROSE='C93D63'; TINT='E9EDF8'
HF=Font(bold=True,color='FFFFFF'); HFILL=PatternFill('solid',fgColor=BLUE); INPUT=PatternFill('solid',fgColor='FFF7DD'); thin=Side(style='thin',color='D9DEE9')
def header(ws,row=1):
    for c in ws[row]:
        if c.value is not None: c.font=HF; c.fill=HFILL; c.alignment=Alignment(wrap_text=True,vertical='top')
def widths(ws,default=16,maxw=48):
    for col in ws.columns:
        L=max((len(str(c.value)) for c in col if c.value is not None),default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width=max(10,min(maxw,L+2))
def cover(wb,title,desc,contents):
    ws=wb.create_sheet('Cover',0)
    ws.sheet_view.showGridLines=False
    try:
        img=XImage(LOGO); img.width,img.height=300,118; ws.add_image(img,'B2')
    except Exception as e: print('logo skipped',e)
    ws['B10']=title; ws['B10'].font=Font(size=24,bold=True,color=BLUE)
    ws['B12']=desc; ws['B12'].font=Font(size=11,color='4B5568'); ws['B12'].alignment=Alignment(wrap_text=True,vertical='top'); ws.merge_cells('B12:H14')
    ws['B16']='Prepared for'; ws['C16']='The Founder, Vytalix'; ws['B17']='Date'; ws['C17']='8 September 2026'; ws['B18']='Version'; ws['C18']='1.0 · Confidential draft'
    for r in (16,17,18): ws[f'B{r}'].font=Font(bold=True,color=BLUE)
    ws['B20']='Contents'; ws['B20'].font=Font(bold=True,size=12,color=BLUE)
    for i,(s,d) in enumerate(contents): ws[f'B{21+i}']=s; ws[f'C{21+i}']=d; ws[f'B{21+i}'].font=Font(bold=True)
    ws[f'B{23+len(contents)}']='Colour code: yellow cells are inputs you can change; blue headers are fixed; every figure is an ESTIMATE unless labelled VERIFIED (see FACTS_BASE).'; ws[f'B{23+len(contents)}'].font=Font(italic=True,color='4B5568')
    ws.column_dimensions['B'].width=18; ws.column_dimensions['C'].width=70
    return ws
def brand_existing(src,dst,title,desc):
    wb=load_workbook(src)
    contents=[(s.title,'') for s in wb.worksheets]
    if 'Cover' in wb.sheetnames: del wb['Cover']
    cover(wb,title,desc,contents)
    for ws in wb.worksheets[1:]:
        if ws.max_row>=1 and any(c.value for c in ws[1]): header(ws)
    wb.save(dst); print('branded',os.path.basename(dst))

# 1) Brand existing workbooks
brand_existing(os.path.join(ROOT,'investors','VYTALIX_INVESTOR_CRM.xlsx'),os.path.join(OUT,'Vytalix_Investor_CRM.xlsx'),'Investor CRM','316 investors scored for fit and contactability, with Top 25, four call lists, pipeline stages and dashboard. Phone numbers only where traced to an official public source.')
brand_existing(os.path.join(ROOT,'finance','Vytalix_Financial_Model_3yr.xlsx'),os.path.join(OUT,'Vytalix_Financial_Model_3yr.xlsx'),'Financial Model (36 months)','Three scenarios (Conservative, Base, High-growth) driven from one assumptions sheet. Change the scenario selector on the Assumptions sheet. Nothing here is an actual; all figures are scenario outputs.')
brand_existing(os.path.join(ROOT,'investors','partners','PARTNER_WAR_ROOM.xlsx'),os.path.join(OUT,'Partner_War_Room.xlsx'),'Partner War Room','76 target organisations and contacts with the job title to approach, a CRM with the eleven-stage partner pipeline, stage rules and monthly KPIs.')

# 2) KPI dashboard and pipeline trackers from CSVs
def from_csv(path,title,desc,sheetname):
    wb=Workbook(); ws=wb.active; ws.title=sheetname
    for r in csv.reader(open(path,encoding='utf-8')): ws.append(r)
    header(ws); widths(ws); ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    return wb,ws
wb,ws=from_csv(os.path.join(ROOT,'operations','KPI_DASHBOARD.csv'),'KPI Dashboard','Commercial, investment, product, marketing and finance KPIs with formulas and month 1–3 targets. Replace EXAMPLE rows with actuals each Friday.','KPIs')
cover(wb,'KPI Dashboard','Commercial, investment, product, marketing and finance KPIs with targets for months 1–3. Replace EXAMPLE rows with actuals every Friday.',[('KPIs','Definitions, formulas, targets, actuals')]); wb.save(os.path.join(OUT,'KPI_Dashboard.xlsx'))
wb,ws=from_csv(os.path.join(ROOT,'operations','PIPELINE_TRACKERS.csv'),'Pipeline Trackers','Investor, partnership, product, revenue and strategic-project pipelines in one sheet.','Pipelines')
cover(wb,'Pipeline Trackers','Investor, partnership, product, revenue and strategic-project pipelines. Filter by the Pipeline column.',[('Pipelines','All five pipelines; filter by type')]); wb.save(os.path.join(OUT,'Pipeline_Trackers.xlsx'))
wb,ws=from_csv(os.path.join(ROOT,'assets','CONTENT_CALENDAR.csv'),'Content Calendar','Twelve weeks of dated content across LinkedIn, website, email, video and events.','Calendar')
cover(wb,'Content Calendar','Twelve weeks from Monday 14 September 2026 across LinkedIn, website, email, video and events; consistent with the marketing engine.',[('Calendar','Dated items with pillar, format, CTA, owner and status')]); wb.save(os.path.join(OUT,'Content_Calendar.xlsx'))

# 3) MaternaLink pricing calculator — three tiers, nine modules
# Rows are tracked by label so formulas can never drift when lines are added.
class Sheet:
    def __init__(self, ws): self.ws=ws; self.r=0; self.rows={}
    def add(self, label=None, value=None, note=None, key=None, fmt=None, inp=False, hdr=False, bold=False):
        self.r+=1
        if label is not None: self.ws[f'A{self.r}']=label
        if value is not None: self.ws[f'B{self.r}']=value
        if note is not None: self.ws[f'C{self.r}']=note
        if key: self.rows[key]=self.r
        if hdr:
            for c in ('A','B','C'):
                self.ws[f'{c}{self.r}'].font=HF; self.ws[f'{c}{self.r}'].fill=HFILL
        if inp: self.ws[f'B{self.r}'].fill=INPUT
        if fmt: self.ws[f'B{self.r}'].number_format=fmt
        if bold:
            self.ws[f'A{self.r}'].font=Font(bold=True,color=BLUE); self.ws[f'B{self.r}'].font=Font(bold=True)
        return self.r
    def b(self,key): return f'B{self.rows[key]}'

wb=Workbook(); S=Sheet(wb.active); wb.active.title='UK Calculator'
S.add('MATERNALINK PRICING CALCULATOR — UK HOSPITAL (change the yellow cells)')
S.ws['A1'].font=Font(bold=True,size=13,color=BLUE)
S.add()
S.add('TIER 1 — every pregnant woman','Value','Notes',hdr=True)
S.add('Births / women enrolled per year',5000,'Drives the site-licence band',key='births',inp=True)
S.add('Tier 1 price per woman per year (£)',22,'Range 16-30',key='t1price',inp=True)
S.add('Site licence (£)',20000,'12k-30k by births',key='site',inp=True)
S.add()
S.add('TIER 2 — women watched on a device','Value','Notes',hdr=True)
S.add('Monitored women per year',800,'Typically 15-25% of births',key='mon',inp=True)
S.add('Price per monitored woman per MONTH (£)',28,'Range 22-38',key='t2price',inp=True)
S.add('Average months monitored',4.2,'16-22 weeks',key='months',inp=True)
S.add('Devices in rotation',300,'Fewer than monitored women; they rotate',key='devices',inp=True)
S.add('Device lease charged per month (£)',6,'Range 4.50-7.00',key='devprice',inp=True)
S.add()
S.add('TIER 3 — the institution','Value','Notes',hdr=True)
S.add('Hospital dashboard (£/year)',28000,'18k-40k by births',key='hosp',inp=True)
S.add('Integration interfaces',1,'',key='intcount',inp=True)
S.add('Integration setup each (£)',25000,'15k-40k, one-off',key='intsetup',inp=True)
S.add('Integration annual each (£)',6000,'4k-8k',key='intann',inp=True)
S.add()
S.add('ADD-ON MODULES','Value','Notes',hdr=True)
S.add('Newborn module per birth (£)',10,'8-14. Excludes jaundice imaging',key='newborn',inp=True)
S.add('Voice diary per woman per year (£)',0,'3-5 once built',key='voice',inp=True)
S.add('Glucose: women with the condition',0,'About 5-10% of the cohort',key='glucount',inp=True)
S.add('Glucose price each per year (£)',25,'20-30, after a regulatory opinion',key='gluprice',inp=True)
S.add('AI engine uplift on Tier 2',0,'0.25-0.35 ONLY after certification',key='ai',inp=True,fmt='0%')
S.add()
S.add('SERVICE AND SETUP','Value','Notes',hdr=True)
S.add('Service plan (0.20 Std / 0.26 Enh / 0.32 Prem)',0.26,'',key='svc',inp=True,fmt='0%')
S.add('Implementation one-off (£)',45000,'25k-75k',key='impl',inp=True)
S.add('Cost to serve one monitored woman per month (£)',6.5,'ESTIMATE — your developer must confirm',key='cost',inp=True)
S.add()
S.add('OUTPUTS','£','',hdr=True)
S.add('Tier 1 revenue',f'={S.b("births")}*{S.b("t1price")}+{S.b("site")}',key='o_t1',fmt='£#,##0')
S.add('Tier 2 revenue',f'={S.b("mon")}*{S.b("t2price")}*{S.b("months")}',key='o_t2',fmt='£#,##0')
S.add('AI engine uplift',f'={S.b("o_t2")}*{S.b("ai")}',key='o_ai',fmt='£#,##0')
S.add('Newborn module',f'={S.b("births")}*{S.b("newborn")}',key='o_nb',fmt='£#,##0')
S.add('Voice diary',f'={S.b("births")}*{S.b("voice")}',key='o_v',fmt='£#,##0')
S.add('Glucose module',f'={S.b("glucount")}*{S.b("gluprice")}',key='o_g',fmt='£#,##0')
S.add('Hospital dashboard',f'={S.b("hosp")}',key='o_h',fmt='£#,##0')
S.add('Integration, annual',f'={S.b("intcount")}*{S.b("intann")}',key='o_i',fmt='£#,##0')
S.add('SOFTWARE SUBTOTAL',f'=SUM({S.b("o_t1")}:{S.b("o_i")})',key='o_sub',fmt='£#,##0',bold=True)
S.add('Service plan',f'={S.b("o_sub")}*{S.b("svc")}',key='o_svc',fmt='£#,##0')
S.add('Device leases (annual)',f'={S.b("devices")}*{S.b("devprice")}*12',key='o_dev',fmt='£#,##0')
S.add('RECURRING ANNUAL TOTAL',f'={S.b("o_sub")}+{S.b("o_svc")}+{S.b("o_dev")}',key='o_rec',fmt='£#,##0',bold=True)
S.add('One-off: implementation + integration build',f'={S.b("impl")}+{S.b("intcount")}*{S.b("intsetup")}',key='o_one',fmt='£#,##0')
S.add('YEAR 1 TOTAL',f'={S.b("o_rec")}+{S.b("o_one")}',key='o_y1',fmt='£#,##0',bold=True)
S.add('YEAR 2+ RUN RATE',f'={S.b("o_rec")}',key='o_rr',fmt='£#,##0',bold=True)
S.add('Year 1 revenue per birth (£)',f'=IF({S.b("births")}>0,{S.b("o_y1")}/{S.b("births")},0)',fmt='£#,##0.00')
S.add('Run-rate revenue per birth (£)',f'=IF({S.b("births")}>0,{S.b("o_rec")}/{S.b("births")},0)',fmt='£#,##0.00')
S.add('Your cost to serve the monitored women',f'={S.b("mon")}*{S.b("months")}*{S.b("cost")}',key='o_cost',fmt='£#,##0')
S.add('Gross margin on Tier 2',f'=IF({S.b("o_t2")}>0,({S.b("o_t2")}-{S.b("o_cost")})/{S.b("o_t2")},0)',fmt='0%')
for col,wd in (('A',52),('B',16),('C',44)): S.ws.column_dimensions[col].width=wd

P=Sheet(wb.create_sheet('Programme Calculator'))
P.add('MATERNALINK PRICING CALCULATOR — AFRICAN PROGRAMME')
P.ws['A1'].font=Font(bold=True,size=13,color=BLUE)
P.add()
P.add('TIER 1 — every woman','Value','Notes',hdr=True)
P.add('Women enrolled per year',100000,'',key='women',inp=True)
P.add('Tier 1 price per woman (£)',None,'Formula applies the published volume tiers',key='t1')
P.ws[P.b('t1')]=f'=IF({P.b("women")}<=5000,9,IF({P.b("women")}<=25000,7,IF({P.b("women")}<=100000,5,4)))'
P.ws[P.b('t1')].fill=PatternFill('solid',fgColor=TINT)
P.add()
P.add('TIER 2 — monitored women','Value','Notes',hdr=True)
P.add('Monitored women',15000,'About 15% of the cohort',key='mon',inp=True)
P.add('Price per monitored woman per month (£)',4,'Range 3-6',key='t2',inp=True)
P.add('Average months monitored',4,'',key='months',inp=True)
P.add()
P.add('MODULES AND INSTITUTION','Value','Notes',hdr=True)
P.add('Newborn module per birth (£)',2,'1.50-3.00',key='nb',inp=True)
P.add('Government / regional dashboard (£)',120000,'60k-250k by population',key='gov',inp=True)
P.add('Facilities with triage',0,'',key='fac',inp=True)
P.add('Triage per facility (£)',2500,'',key='facprice',inp=True)
P.add()
P.add('SERVICE AND SETUP','Value','Notes',hdr=True)
P.add('Service plan (0.20 / 0.28 / 0.35)',0.28,'',key='svc',inp=True,fmt='0%')
P.add('Training cohorts',30,'',key='coh',inp=True)
P.add('Training per cohort (£)',1800,'',key='cohprice',inp=True)
P.add('Implementation one-off (£)',150000,'40k-200k by size',key='impl',inp=True)
P.add()
P.add('OUTPUTS','£','',hdr=True)
P.add('Tier 1 revenue',f'={P.b("women")}*{P.b("t1")}',key='o1',fmt='£#,##0')
P.add('Tier 2 revenue',f'={P.b("mon")}*{P.b("t2")}*{P.b("months")}',key='o2',fmt='£#,##0')
P.add('Newborn module',f'={P.b("women")}*{P.b("nb")}',key='o3',fmt='£#,##0')
P.add('Government dashboard',f'={P.b("gov")}',key='o4',fmt='£#,##0')
P.add('Triage',f'={P.b("fac")}*{P.b("facprice")}',key='o5',fmt='£#,##0')
P.add('SOFTWARE SUBTOTAL',f'=SUM({P.b("o1")}:{P.b("o5")})',key='osub',fmt='£#,##0',bold=True)
P.add('Service plan',f'=MAX({P.b("osub")}*{P.b("svc")},12000)','Minimum £12,000',key='osvc',fmt='£#,##0')
P.add('Training',f'={P.b("coh")}*{P.b("cohprice")}','Recurs: new staff every year',key='otr',fmt='£#,##0')
P.add('RECURRING ANNUAL TOTAL',f'={P.b("osub")}+{P.b("osvc")}+{P.b("otr")}',key='orec',fmt='£#,##0',bold=True)
P.add('YEAR 1 TOTAL',f'={P.b("orec")}+{P.b("impl")}',key='oy1',fmt='£#,##0',bold=True)
P.add('YEAR 2+ RUN RATE',f'={P.b("orec")}',key='orr',fmt='£#,##0',bold=True)
P.add('Year 1 cost per woman (£)',f'=IF({P.b("women")}>0,{P.b("oy1")}/{P.b("women")},0)',fmt='£#,##0.00')
P.add('Run-rate cost per woman (£)',f'=IF({P.b("women")}>0,{P.b("orr")}/{P.b("women")},0)',fmt='£#,##0.00')
for col,wd in (('A',52),('B',16),('C',44)): P.ws.column_dimensions[col].width=wd

ws3=wb.create_sheet('Price Book')
for r in [['Tier / module','What it covers','UK price','Programme price','Sell it yet?'],
['TIER 1 — everyone','App, appointments, education, symptoms, mood, fetal movement log','£16-30 per woman per year + £12k-30k site licence','£4-9 per woman per year','Yes'],
['TIER 2 — watched closely','Connected BP monitor, active surveillance, escalation','£22-38 per woman per MONTH','£3-6 per woman per month','Yes'],
['TIER 3 — hospital dashboard','Population view, response times, analytics','£18k-40k per site per year','Included in programme','Yes'],
['TIER 3 — government dashboard','Regional and national intelligence, de-identified data','n/a','£60k-250k per year','Yes'],
['Clinician dashboard','Patient list, red/amber/green, graphs, timeline, escalation','Included in Tiers 1 and 2','Included','Yes, if the rules belong to the hospital'],
['Newborn module','Feeding, temperature, breathing, observations','£8-14 per birth','£1.50-3 per birth','Yes, without jaundice imaging'],
['Voice diary','Storage and playback only, no analysis','£3-5 per woman per year','£1-2','Yes'],
['Glucose diary','Self-reported readings against clinician targets','£20-30 per woman with the condition','£6-10','After a regulatory opinion'],
['Integration','EPR/FHIR, secure API, access control, audit trail','£15k-40k setup + £4k-8k a year','£8k-25k + £3k','Yes'],
['AI ENGINE','Maternal Instability Score, risk stratification, deterioration detection','Add 25-35% to Tier 2','Add 20%','NO — needs a medical device licence'],
['Jaundice photo screening','Image-assisted newborn screening','Not priced','Not priced','NO — a device, and heavily patented already'],
['Devices','Validated BP monitor, thermometer','Lease £4.50-7 per month, or cost + 15%','Lease £2-4 per month','Yes, resell certified kit'],
['Implementation','Setup, configuration, training, safety paperwork','£25k-75k per site','£40k-200k per programme','Yes'],
['Service plan','Support, releases, patching, safety maintenance','20% / 26% / 32% of licence','20% / 28% / 35%','Yes'],
['Pilot','12 weeks, one site, evaluation report','£40k-90k','On scope','Yes']]: ws3.append(r)
header(ws3); widths(ws3,maxw=46)

ws4=wb.create_sheet('Packages')
for r in [['Package','What is in it','5,000-birth hospital, year 1','Run rate'],
['Essential','Tier 1 for all, Tier 2 for 500 women, clinician dashboard, service plan','About £250,000','About £205,000'],
['Complete','Essential plus hospital dashboard, newborn module, one integration','About £390,000','About £320,000'],
['Full ecosystem','Complete plus 1,000 monitored, voice diary, glucose, premium support','About £490,000','About £410,000'],
[],
['Rule','Detail'],
['Pilot first','12 weeks fixed price, half the fee credited against year one'],
['Minimum contract','£40,000 in the UK, £25,000 for a programme'],
['No discounting below the floor','Remove modules instead of cutting the price'],
['Devices always separate','Never inside the software price'],
['Nothing unlicensed on a price list','The AI engine and jaundice imaging stay off until certified']]: ws4.append(r)
header(ws4); widths(ws4,maxw=60)

cover(wb,'MaternaLink Pricing Calculator','Three tiers across nine modules. Yellow cells are inputs you can change. Two modules are marked as not sellable until they hold a medical device licence. Every price is an ESTIMATE until tested with three buyers per segment.',[('UK Calculator','A UK hospital, tier by tier, with cost to serve and margin'),('Programme Calculator','An African state or NGO programme, tiered by volume'),('Price Book','Every module, both markets, and whether you can sell it yet'),('Packages','The three packages to offer, and the rules you do not break')])
wb.save(os.path.join(OUT,'MaternaLink_Pricing_Calculator.xlsx')); print('pricing calculator ok')

# 4) Investor readiness scorecard
wb=Workbook(); ws=wb.active; ws.title='Scorecard'
ws.append(['#','Readiness item','Today (0/1)','Target','How to close','Owner','Due','Status','Notes'])
items=[('Legal entity, bank, accountant',0,1,'Name clearance, incorporate, bank, accountant, SEIS/EIS advance assurance','Founder','Day 10'),('MaternaLink IP owned or licensed in writing',0,1,'Heads of terms then IP deed with collaborator; contributor assignments','Founder + solicitor','Day 30'),('No unvalidated clinical claims public',0,1,'Retire 2025 deck; relabel prototype rules; auth; remove photos','Founder','Day 14'),('Name and trademark clearance',0,1,'Knock-out searches; file UK TM classes 9, 41, 42, 44','Founder + attorney','Day 20'),('First revenue',0,1,'Advisory packages; 40 conversations; first SOW','Founder','Day 45'),('Pilot partner LOI',0,1,'15 discovery interviews; pilot proposal','Founder','Day 60'),('Advisory board (2 confirmed)',0,1,'Named asks; advisory agreements','Founder','Day 60'),('Product evidence (audit, PRD, intended use, hazard log)',0,1,'Code audit; clinician sign-off','CTO/CPO','Day 30'),('Financial model with actuals',1,1,'Update with months 1–3 actuals','CFO role','Day 60'),('Data room',0,1,'Index in doc 23 §2','Founder','Day 45'),('Investor list and outreach system',1,1,'Verify Call List 1 phones','Founder','Day 55'),('Compliance basics (ICO, Cyber Essentials, insurance)',0,1,'See doc 15','Founder','Day 30'),('Founder commitment statement and time split',0,1,'Written statement','Founder','Day 7'),('Proof of demand (15 interviews)',0,1,'Discovery interviews','Founder','Day 60'),('Non-dilutive application submitted',0,1,'Innovate UK Smart / SBRI Healthcare','Founder','Day 75')]
for i,(a,t,g,h,o,d) in enumerate(items,1): ws.append([i,a,t,g,h,o,d,'Not started',''])
n=len(items)+1
ws.append([]); ws.append(['','READINESS SCORE /10',f'=ROUND(SUM(C2:C{n})/COUNT(C2:C{n})*10,1)','', 'Rule: no first investor meeting until items 1–4 are done and two of 5–7 are done'])
ws[f'B{n+2}'].font=Font(bold=True,color=BLUE); ws[f'C{n+2}'].font=Font(bold=True,size=12)
header(ws); widths(ws,maxw=60); ws.freeze_panes='A2'
for r in range(2,n+1): ws[f'C{r}'].fill=INPUT; ws[f'H{r}'].fill=INPUT
cover(wb,'Investor Readiness Scorecard','Fifteen readiness items with owners and due dates. Update the Today column (0 or 1) weekly; the score recalculates. Rule: no first investor meeting until items 1–4 are done and two of 5–7 are done.',[('Scorecard','Items, targets, owners, due dates, live score')])
wb.save(os.path.join(OUT,'Investor_Readiness_Scorecard.xlsx')); print('scorecard ok')
print('xlsx done')
