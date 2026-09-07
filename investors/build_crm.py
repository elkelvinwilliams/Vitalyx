#!/usr/bin/env python3
"""Merge raw research CSVs into the Vytalix Investor CRM (CSV + XLSX + Markdown outputs).

Usage: python3 investors/build_crm.py
Inputs : investors/raw/*.csv (SCHEMA.md columns)
Outputs: investors/VYTALIX_INVESTOR_CRM.csv, investors/VYTALIX_INVESTOR_CRM.xlsx,
         investors/TOP_25_INVESTORS.md, investors/CALL_LISTS.md, investors/OUTREACH_TOP25.md,
         investors/INVESTOR_DASHBOARD.md
Scoring is deterministic from the sub-scores captured during research (see SCHEMA.md rubric).
"""
import csv, glob, os, re, sys, datetime
from collections import OrderedDict, Counter

ROOT = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(ROOT, "raw")
TODAY = "2026-09-07"

SCHEMA = [l.strip() for l in open(os.path.join(RAW, "SCHEMA.md")) if l.startswith("investor_name,")][0].split(",")
FIT = ["sector_fit_20","stage_fit_15","ticket_fit_15","geo_fit_10","strategic_fit_15","maternalink_fit_10","accessibility_5","network_value_10"]
FITMAX = {"sector_fit_20":20,"stage_fit_15":15,"ticket_fit_15":15,"geo_fit_10":10,"strategic_fit_15":15,"maternalink_fit_10":10,"accessibility_5":5,"network_value_10":10}
CON = ["c_phone_3","c_direct_route_2","c_named_dm_1","c_email_1","c_linkedin_1","c_warm_intro_1","c_info_quality_1"]
CONMAX = {"c_phone_3":3,"c_direct_route_2":2,"c_named_dm_1":1,"c_email_1":1,"c_linkedin_1":1,"c_warm_intro_1":1,"c_info_quality_1":1}
PIPELINE = ["Identified","Researched","Qualified","Priority A","Outreach Ready","Contacted","Called","Email Sent","Responded","Meeting Requested","Meeting Booked","Meeting Completed","Due Diligence","Investment Discussion","Term Sheet","Closed","Passed","Future Opportunity"]

def norm(name):
    n = name.lower()
    n = re.sub(r"\(.*?\)", "", n)
    n = re.sub(r"\b(ltd|limited|llp|lp|plc|inc|gmbh|sa|ag|the|ventures?|capital|partners?|fund|funds|group|management|investments?|vc)\b", " ", n)
    return re.sub(r"[^a-z0-9]+", "", n)

def clip(v, mx):
    try:
        x = int(round(float(str(v).strip() or 0)))
    except ValueError:
        x = 0
    return max(0, min(mx, x))

def is_phone_verified(r):
    return str(r.get("phone_verification_status","")).upper().startswith("VERIFIED") and any(r.get(k,"").strip() for k in ("phone_investment_team","phone_office","phone_investor_relations"))

def best_phone(r):
    for k in ("phone_investment_team","phone_office","phone_investor_relations"):
        if r.get(k,"").strip():
            return r[k].strip()
    return "PHONE NOT PUBLICLY AVAILABLE"

def load():
    rows, seen, dupes, bad = [], {}, [], []
    for path in sorted(glob.glob(os.path.join(RAW, "*.csv"))):
        cat = os.path.basename(path)
        with open(path, newline="", encoding="utf-8") as f:
            rd = csv.DictReader(f)
            if [h.strip() for h in rd.fieldnames or []] != SCHEMA:
                bad.append((cat, rd.fieldnames)); continue
            for r in rd:
                r = {k: (v or "").strip() for k, v in r.items()}
                if not r.get("investor_name"): continue
                key = norm(r["investor_name"])
                if key in seen:
                    dupes.append((r["investor_name"], seen[key])); continue
                seen[key] = r["investor_name"]
                r["_category_file"] = cat
                rows.append(r)
    return rows, dupes, bad

def score(r):
    fit = sum(clip(r.get(k), FITMAX[k]) for k in FIT)
    # contactability: phone verified is the anchor; cap the phone sub-score by evidence
    if not is_phone_verified(r):
        r["c_phone_3"] = "0"
    con = sum(clip(r.get(k), CONMAX[k]) for k in CON)
    overall = round(fit * 0.85 + con * 1.5)
    tier = "A" if fit >= 80 else "B" if fit >= 65 else "C" if fit >= 50 else "D"
    prob = "High" if (fit >= 75 and con >= 6) else "Medium" if (fit >= 60 and con >= 4) else "Low"
    return fit, con, overall, tier, prob

def next_action(r, tier, con):
    if tier == "D": return "Hold. Re-score only if new information emerges."
    pv = is_phone_verified(r)
    dm = r.get("decision_maker_name","") not in ("", "NOT PUBLICLY IDENTIFIED")
    if tier == "A" and pv: return "Call investment team using verified number; follow with personalised email same day."
    if tier == "A": return "Verify phone via official contact page; send personalised email/LinkedIn to " + (r.get("decision_maker_name") if dm else "investment team") + "; seek warm intro."
    if tier == "B" and pv: return "Call to qualify fit and process; log outcome; send one-pager."
    if tier == "B": return "Email/LinkedIn introduction; request 20-minute call; verify phone."
    return "Monitor; add to quarterly update list; approach after first pilot/revenue."

def build():
    rows, dupes, bad = load()
    for cat, hdr in bad:
        print(f"WARNING: header mismatch in {cat}: {hdr}", file=sys.stderr)
    for r in rows:
        r["fit_score_100"], r["contactability_10"], r["overall_score"], r["tier"], r["investment_probability"] = score(r)
    rows.sort(key=lambda r: (-r["overall_score"], -r["fit_score_100"], r["investor_name"]))
    for i, r in enumerate(rows, 1):
        r["investor_id"] = f"VX-{i:03d}"
        r["priority"] = {"A":"A-Tier (80–100): immediate personalised outreach","B":"B-Tier (65–79): targeted relationship development","C":"C-Tier (50–64): maintain and monitor","D":"Below 50: do not prioritise"}[r["tier"]]
        r["crm_status"] = "Researched"
        r["next_action"] = next_action(r, r["tier"], r["contactability_10"])
        r["best_phone"] = best_phone(r)
    return rows, dupes

OUT_COLS = OrderedDict([
    ("Investor ID","investor_id"),("Investor / Organisation","investor_name"),("Investor Type","investor_type"),("Subcategory","subcategory"),
    ("Country","country"),("City","city"),("Official Website","website"),("LinkedIn","linkedin_url"),
    ("Professional Phone Number (best)","best_phone"),("Investment Contact Phone","phone_investment_team"),("General Office Phone","phone_office"),("Investor Relations Phone","phone_investor_relations"),
    ("Phone Verification Status","phone_verification_status"),("Phone Source","phone_source_url"),
    ("Investment Email","email_investment"),("General Email","email_general"),("Email Source","email_source_url"),
    ("Investment Focus","investment_focus"),("Healthcare Focus","healthcare_focus"),("Technology Focus","technology_focus"),("Geography Focus","geography_focus"),
    ("Investment Stage","investment_stage"),("Typical Ticket","typical_ticket"),("Fund Size / Assets","fund_size_or_aum"),("Relevant Portfolio","relevant_portfolio"),
    ("Vytalix Fit (Sector /20)","sector_fit_20"),("Stage Fit /15","stage_fit_15"),("Ticket Fit /15","ticket_fit_15"),("Geographic Fit /10","geo_fit_10"),("Strategic Value /15","strategic_fit_15"),("MaternaLink Fit /10","maternalink_fit_10"),("Accessibility /5","accessibility_5"),("Network Value /10","network_value_10"),
    ("VYTALIX INVESTOR FIT SCORE /100","fit_score_100"),("CONTACTABILITY SCORE /10","contactability_10"),("Investment Probability","investment_probability"),("Overall Score","overall_score"),("Tier","tier"),("Priority","priority"),
    ("Relevant Decision Maker","decision_maker_name"),("Decision Maker Role","decision_maker_role"),("Decision Maker Source","decision_maker_source_url"),
    ("Contact Route","contact_route"),("Warm Intro Possibility","warm_intro_possibility"),("Source","source_urls"),("Last Verified","last_verified"),
    ("CRM Status","crm_status"),("Next Action","next_action"),("Why Investor for Vytalix","why_investor_for_vytalix"),("Likely Objection","likely_objection"),("Recommended Pitch Angle","pitch_angle"),("Notes","notes"),("Research Batch","_category_file"),
])

def write_csv(rows):
    p = os.path.join(ROOT, "VYTALIX_INVESTOR_CRM.csv")
    with open(p, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL)
        w.writerow(OUT_COLS.keys())
        for r in rows:
            w.writerow([r.get(k, "") for k in OUT_COLS.values()])
    return p

def call_lists(rows):
    active = [r for r in rows if r["tier"] != "D"]
    return {"CALL LIST 1 — Top 25": active[:25], "CALL LIST 2 — Next 50": active[25:75], "CALL LIST 3 — Next 75": active[75:150], "CALL LIST 4 — Remaining prospects": active[150:]}

def dashboard(rows):
    c = Counter(r["tier"] for r in rows)
    pv = sum(1 for r in rows if is_phone_verified(r))
    ev = sum(1 for r in rows if (r.get("email_investment") or r.get("email_general")) and "NOT PUBLICLY" not in (r.get("email_investment","")+r.get("email_general","")).upper())
    dm = sum(1 for r in rows if r.get("decision_maker_name") and r["decision_maker_name"] != "NOT PUBLICLY IDENTIFIED")
    types = Counter(r["investor_type"] for r in rows)
    countries = Counter(r["country"] for r in rows)
    return c, pv, ev, dm, types, countries

def write_md(rows, dupes):
    c, pv, ev, dm, types, countries = dashboard(rows)
    n = len(rows)
    # Dashboard
    with open(os.path.join(ROOT, "INVESTOR_DASHBOARD.md"), "w", encoding="utf-8") as f:
        f.write(f"# VYTALIX INVESTOR DASHBOARD\n\nGenerated {TODAY} from `VYTALIX_INVESTOR_CRM.csv`. Activity counters are zero because outreach has not started; update `CRM Status` in the CRM and re-run `build_crm.py`.\n\n")
        f.write("| Metric | Value |\n|---|---|\n")
        f.write(f"| TOTAL INVESTORS | {n} |\n| A-TIER (fit 80–100) | {c['A']} |\n| B-TIER (65–79) | {c['B']} |\n| C-TIER (50–64) | {c['C']} |\n| BELOW 50 | {c['D']} |\n")
        f.write(f"| PHONE NUMBERS VERIFIED (public, sourced) | {pv} |\n| PHONE NOT PUBLICLY AVAILABLE | {n-pv} |\n| EMAILS VERIFIED (public, sourced) | {ev} |\n| DECISION MAKERS IDENTIFIED (public role) | {dm} |\n")
        st = Counter(r["crm_status"] for r in rows)
        contacted = sum(v for k,v in st.items() if k in PIPELINE[5:])
        f.write(f"| INVESTORS CONTACTED | {contacted} |\n| CALLS MADE | {st.get('Called',0)} |\n| RESPONSES | {sum(v for k,v in st.items() if k in PIPELINE[8:])} |\n| MEETINGS | {sum(v for k,v in st.items() if k in PIPELINE[10:])} |\n| ACTIVE INVESTMENT DISCUSSIONS | {st.get('Investment Discussion',0)+st.get('Due Diligence',0)+st.get('Term Sheet',0)} |\n| CAPITAL RAISED | £0 (pre-formation) |\n| CONVERSION RATE (contacted → meeting) | n/a (no outreach yet) |\n\n")
        f.write("## By investor type\n\n| Type | Count | A | B | C |\n|---|---|---|---|---|\n")
        for t, k in types.most_common():
            f.write(f"| {t} | {k} | {sum(1 for r in rows if r['investor_type']==t and r['tier']=='A')} | {sum(1 for r in rows if r['investor_type']==t and r['tier']=='B')} | {sum(1 for r in rows if r['investor_type']==t and r['tier']=='C')} |\n")
        f.write("\n## By country (top 15)\n\n| Country | Count |\n|---|---|\n")
        for t, k in countries.most_common(15): f.write(f"| {t} | {k} |\n")
        f.write("\n## Pipeline stages (all records start at Researched)\n\n" + " → ".join(PIPELINE) + "\n")
        if dupes:
            f.write(f"\n## De-duplication log\n\n{len(dupes)} duplicate records removed during merge: " + "; ".join(f"{a} (kept as {b})" for a,b in dupes) + "\n")
    # Top 25
    top = [r for r in rows if r["tier"] != "D"][:25]
    with open(os.path.join(ROOT, "TOP_25_INVESTORS.md"), "w", encoding="utf-8") as f:
        f.write("# TOP 25 INVESTORS — first approach list\n\nRanked by Overall Score (0.85 × Fit /100 + 1.5 × Contactability /10). Phone numbers appear only where a public source was captured; otherwise PHONE NOT PUBLICLY AVAILABLE and the contact route is given.\n\n")
        f.write("| Rank | Investor | Type | Country | Fit | Contact. | Phone (best) | Route |\n|---|---|---|---|---|---|---|---|\n")
        for i, r in enumerate(top, 1):
            f.write(f"| {i} | {r['investor_name']} | {r['investor_type']} | {r['country']} | {r['fit_score_100']} | {r['contactability_10']} | {r['best_phone']} | {r['contact_route'][:80]} |\n")
        for i, r in enumerate(top, 1):
            f.write(f"\n## {i}. {r['investor_name']} ({r['investor_id']})\n\n")
            f.write(f"1. **Investor:** {r['investor_name']} — {r['investor_type']} ({r['subcategory']}), {r['city']}, {r['country']} — {r['website']}\n")
            f.write(f"2. **Why they fit Vytalix:** {r['why_investor_for_vytalix']}\n")
            f.write(f"3. **Why they fit MaternaLink:** MaternaLink fit {r['maternalink_fit_10']}/10 — {r['healthcare_focus'] or r['investment_focus']}\n")
            f.write(f"4. **Investment stage:** {r['investment_stage']}\n5. **Typical ticket:** {r['typical_ticket']}\n")
            f.write(f"6. **Phone number:** {r['best_phone']} ({r['phone_verification_status']}{'; source: '+r['phone_source_url'] if r['phone_source_url'] else ''})\n")
            f.write(f"7. **Email / contact route:** {r['email_investment'] or r['email_general'] or 'NOT PUBLICLY AVAILABLE'} — {r['contact_route']}\n")
            f.write(f"8. **Relevant decision maker:** {r['decision_maker_name']} ({r['decision_maker_role']})\n")
            f.write(f"9. **Best introduction route:** {r['warm_intro_possibility']}\n")
            f.write(f"10. **Relevant previous investment:** {r['relevant_portfolio']}\n")
            f.write(f"11. **Strategic value:** {r['strategic_fit_15']}/15 strategic, {r['network_value_10']}/10 network — {r['notes'][:220]}\n")
            f.write(f"12. **Likely objection:** {r['likely_objection']}\n13. **Recommended pitch angle:** {r['pitch_angle']}\n")
            f.write(f"14. **Priority score:** Fit {r['fit_score_100']}/100 · Contactability {r['contactability_10']}/10 · Overall {r['overall_score']} · Tier {r['tier']}\n")
            f.write(f"15. **Next action:** {r['next_action']}\n")
    # Call lists
    with open(os.path.join(ROOT, "CALL_LISTS.md"), "w", encoding="utf-8") as f:
        f.write("# INVESTOR CALLING CAMPAIGN — CALL LISTS\n\nRules: call only after the entity is incorporated and the MaternaLink IP position is agreed (see `/docs/00_MASTER_AUDIT.md`). Where phone is NOT PUBLICLY AVAILABLE, the first action is to verify a number via the official contact page or to use the stated route. Never cold-call a personal mobile. Log every attempt in `CRM Status`.\n\n")
        for title, lst in call_lists(rows).items():
            f.write(f"## {title} ({len(lst)} investors)\n\n| # | Investor | Phone | Contact person | Type | Fit | Why calling | Opening line | Pitch angle | Objective | Follow-up | CRM status |\n|---|---|---|---|---|---|---|---|---|---|---|---|\n")
            for i, r in enumerate(lst, 1):
                dm = r['decision_maker_name'] if r['decision_maker_name'] and r['decision_maker_name']!='NOT PUBLICLY IDENTIFIED' else 'Investment team'
                opening = f"Good morning, this is [Founder], founder of Vytalix, a UK health-technology group. I'm calling because {r['investor_name']} backs {(r['relevant_portfolio'].split(';')[0].split(',')[0] or 'early-stage health-tech')[:60]} and I'd value 20 minutes with {dm} on a maternal-health platform we're developing. Who is the right person to speak to?"
                f.write(f"| {i} | {r['investor_name']} | {r['best_phone']} | {dm} | {r['investor_type']} | {r['fit_score_100']} | {r['why_investor_for_vytalix'][:120]} | {opening} | {r['pitch_angle'][:120]} | Book a 20-min intro call; confirm stage/ticket fit; identify process | Email one-pager same day; follow-up call day 5; LinkedIn day 10 | {r['crm_status']} |\n")
            f.write("\n")
    # Personalised outreach for top 25
    with open(os.path.join(ROOT, "OUTREACH_TOP25.md"), "w", encoding="utf-8") as f:
        f.write("# PERSONALISED OUTREACH — TOP 25\n\nGenerated from CRM fields; edit for voice before sending. Replace [Founder], [company number] and [link] once incorporated. Never send before the IP position and entity are in place. See `OUTREACH_ENGINE.md` for the full sequence, meeting structure and close.\n")
        for i, r in enumerate(top, 1):
            dm = r['decision_maker_name'] if r['decision_maker_name'] and r['decision_maker_name']!='NOT PUBLICLY IDENTIFIED' else None
            first = dm.split()[0] if dm else "there"
            port = (r['relevant_portfolio'].split(';')[0].split(',')[0]).strip() or "your health-tech portfolio"
            f.write(f"\n---\n\n## {i}. {r['investor_name']}\n\n**PHONE OPENING**\n\n\"Hello, my name is [Founder], founder of Vytalix, a UK health-technology group building MaternaLink, a maternal-health coordination platform in development for the UK and Nigeria. I'm calling {r['investor_name']} because of your work with {port}. I'm not asking for a decision today; I'd like 20 minutes with {dm or 'the investment team'} to test whether this is a fit for your {r['investment_stage'] or 'early-stage'} mandate. Is {dm or 'the investment team'} the right route, and what's the best way to get a slot?\"\n\n")
            f.write(f"**EMAIL**\n\nSubject: Maternal-health platform (UK + Nigeria) — 20 minutes with {r['investor_name']}?\n\nDear {first},\n\nI'm the founder of Vytalix, a UK health-technology group. We're developing MaternaLink, a maternal care-coordination and communication platform for health systems in the UK and Nigeria, alongside an advisory practice that funds our early work.\n\nI'm writing to you specifically because {r['investor_name']} backs {port}, and because {r['why_investor_for_vytalix'][0].lower()+r['why_investor_for_vytalix'][1:]}\n\nWhere we are, honestly: pre-seed, founder-led, a working prototype, an open conversation with a Nigerian state ministry, and a UK pilot programme in design. We are raising a small SEIS/EIS-eligible pre-seed round to fund the pilot and clinical advisory work [confirm once assured]. {r['pitch_angle']}\n\nCould I have 20 minutes to test the fit against your mandate? I'll send a one-pager ahead of the call.\n\nKind regards,\n[Founder]\nFounder, Vytalix — Technology for Life\n[phone] · [email] · [link]\n\n")
            f.write(f"**LINKEDIN**\n\nConnection note (≤300 chars): \"Hi {first}, founder of Vytalix here (UK health-tech; maternal-health platform in development for UK + Nigeria). Following {r['investor_name']}'s work with {port}. Would value connecting.\"\n\nMessage after connection: \"Thanks for connecting, {first}. Quick context: we're pre-seed, building MaternaLink with a services arm funding the early work. {r['pitch_angle'][:160]} Open to a 20-minute call in the next two weeks? Happy to send a one-pager first.\"\n\n")
            f.write(f"**FOLLOW-UP #1 (day 4):** \"Sharing the one-pager as promised [link]. The specific question I'd like your view on: {r['likely_objection'][:140]} — I'd rather hear it early.\"\n\n**FOLLOW-UP #2 (day 10):** \"One update since I wrote: [milestone, e.g., advisory board member confirmed / pilot LOI]. If timing is wrong, could you point me to the right person at {r['investor_name']} or a fund you'd suggest?\"\n\n**FOLLOW-UP #3 (day 21):** \"Closing the loop. I'll add you to our quarterly investor update unless you'd prefer not; next update covers pilot progress and the first revenue quarter. Thanks for your time.\"\n\n")
            f.write(f"**MEETING (20–30 min):** 1) Their mandate and process (5) · 2) Problem and why now (3) · 3) What exists today, no inflation (4) · 4) MaternaLink v1 and evidence plan (5) · 5) Business model: services + platform (3) · 6) Round, use of funds, milestones (3) · 7) Their objection: '{r['likely_objection'][:100]}' answered directly (3) · 8) Agree next step (2).\n\n**CLOSE:** Ask: \"Based on what you've seen, is this within mandate? If yes, what would you need to see to lead or co-invest at pre-seed, and by when?\" Convert to: data-room access → second call with a partner → term discussion. Log stage in CRM.\n")
    print(f"CRM rows: {n} | A {c['A']} B {c['B']} C {c['C']} D {c['D']} | phones verified {pv} | emails {ev} | decision makers {dm} | dupes removed {len(dupes)}")

def write_xlsx(rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not available; skipped xlsx", file=sys.stderr); return
    wb = Workbook()
    ws = wb.active; ws.title = "03 MASTER CRM"
    hdr = list(OUT_COLS.keys())
    ws.append(hdr)
    for r in rows: ws.append([r.get(k, "") for k in OUT_COLS.values()])
    fill = {"A":"FFF2CC","B":"E2EFDA","C":"DDEBF7","D":"EDEDED"}
    tcol = hdr.index("Tier")+1
    for row in ws.iter_rows(min_row=2):
        t = row[tcol-1].value
        for cell in row: cell.fill = PatternFill("solid", fgColor=fill.get(t,"FFFFFF"))
    for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="0B3C5D"); cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(hdr, 1): ws.column_dimensions[get_column_letter(i)].width = 14 if len(h) < 16 else 28
    # Dashboard
    c, pv, ev, dm, types, countries = dashboard(rows)
    d = wb.create_sheet("01 DASHBOARD", 0)
    d.append(["VYTALIX INVESTOR DASHBOARD", f"Generated {TODAY}"])
    n = len(rows)
    for k, v in [("TOTAL INVESTORS", n),("A-TIER", c['A']),("B-TIER", c['B']),("C-TIER", c['C']),("BELOW 50", c['D']),("PHONE NUMBERS VERIFIED", pv),("EMAILS VERIFIED", ev),("DECISION MAKERS IDENTIFIED", dm),("INVESTORS CONTACTED", 0),("CALLS MADE",0),("RESPONSES",0),("MEETINGS",0),("ACTIVE INVESTMENT DISCUSSIONS",0),("CAPITAL RAISED","£0"),("CONVERSION RATE","n/a")]:
        d.append([k, v])
    d.append([]); d.append(["By type"]) 
    for t, k in types.most_common(): d.append([t, k])
    d.column_dimensions["A"].width = 36; d.column_dimensions["B"].width = 18
    d["A1"].font = Font(bold=True, size=14)
    # Top 25 + call lists
    t25 = wb.create_sheet("02 TOP 25")
    t25.append(["Rank","Investor ID","Investor","Type","Country","Fit /100","Contactability /10","Overall","Phone (best)","Decision maker","Role","Contact route","Pitch angle","Likely objection","Next action"])
    for i, r in enumerate([x for x in rows if x["tier"]!="D"][:25], 1):
        t25.append([i, r["investor_id"], r["investor_name"], r["investor_type"], r["country"], r["fit_score_100"], r["contactability_10"], r["overall_score"], r["best_phone"], r["decision_maker_name"], r["decision_maker_role"], r["contact_route"], r["pitch_angle"], r["likely_objection"], r["next_action"]])
    for title, lst in call_lists(rows).items():
        s = wb.create_sheet(("04 " + title.split(" — ")[0]).replace("CALL LIST","CALL LIST"))
        s.append([title]); s.append(["#","Investor ID","Investor","Phone","Contact person","Investor type","Fit","Why calling","Pitch angle","Objective","Follow-up","CRM status","Outcome / notes"])
        for i, r in enumerate(lst, 1):
            dm_ = r['decision_maker_name'] if r['decision_maker_name'] and r['decision_maker_name']!='NOT PUBLICLY IDENTIFIED' else 'Investment team'
            s.append([i, r["investor_id"], r["investor_name"], r["best_phone"], dm_, r["investor_type"], r["fit_score_100"], r["why_investor_for_vytalix"], r["pitch_angle"], "Book 20-min intro; confirm fit; identify process", "Email one-pager same day; call day 5; LinkedIn day 10", r["crm_status"], ""])
    dd = wb.create_sheet("05 PIPELINE & DICTIONARY")
    dd.append(["Pipeline stages"]); dd.append(PIPELINE); dd.append([])
    dd.append(["Fit score /100 = Sector 20 + Stage 15 + Ticket 15 + Geo 10 + Strategic 15 + MaternaLink 10 + Accessibility 5 + Network 10"])
    dd.append(["Contactability /10 = phone verified 3 + direct route 2 + named decision maker 1 + email 1 + LinkedIn 1 + warm intro 1 + info quality 1"])
    dd.append(["Overall = 0.85 × Fit + 1.5 × Contactability. Tiers: A 80–100, B 65–79, C 50–64, below 50 not prioritised (by Fit)."])
    dd.append(["Phone numbers: only public numbers captured from official sources during research; PHONE NOT PUBLICLY AVAILABLE otherwise. Never fabricated. Personal mobiles never recorded."])
    wb.save(os.path.join(ROOT, "VYTALIX_INVESTOR_CRM.xlsx"))

if __name__ == "__main__":
    rows, dupes = build()
    write_csv(rows); write_md(rows, dupes); write_xlsx(rows)
